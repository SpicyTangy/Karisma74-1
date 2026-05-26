"""L'Excel ha due fogli: 'Karisma' (documenti del fornitore, fornitore_ok=True)
e 'Scartati' (non pertinenti o in errore). Le cifre dei non pertinenti (es. MSC)
non devono mai finire nel foglio Karisma."""
from openpyxl import load_workbook

from excel_writer import write_excel


def _rec(file, fornitore_ok, campi=None, status="x"):
    return {"file": file, "fornitore_ok": fornitore_ok, "status": status, "campi": campi or {}}


def test_due_fogli_separano_karisma_dagli_scartati(tmp_path):
    records = [
        _rec("karisma1.pdf", True, {"totale": "100"}, "classificato"),
        _rec("msc.pdf", False, {"totale": "2896.30"}, "non_pertinente"),
        _rec("errore.pdf", None, {}, "errore"),
    ]
    out = tmp_path / "r.xlsx"
    write_excel(records, str(out))

    wb = load_workbook(out)
    assert wb.sheetnames == ["Karisma", "Scartati"]
    karisma = [row[0] for row in wb["Karisma"].iter_rows(min_row=2, values_only=True)]
    scartati = [row[0] for row in wb["Scartati"].iter_rows(min_row=2, values_only=True)]
    assert karisma == ["karisma1.pdf"]
    assert set(scartati) == {"msc.pdf", "errore.pdf"}


def test_cifre_msc_non_finiscono_nel_foglio_karisma(tmp_path):
    records = [_rec("msc.pdf", False, {"totale": "2896.30", "commissioni": "316.44"}, "non_pertinente")]
    out = tmp_path / "r.xlsx"
    write_excel(records, str(out))

    wb = load_workbook(out)
    assert list(wb["Karisma"].iter_rows(min_row=2, values_only=True)) == []   # nessuna riga MSC
    assert len(list(wb["Scartati"].iter_rows(min_row=2, values_only=True))) == 1
