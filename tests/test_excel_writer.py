from openpyxl import load_workbook
from excel_writer import FIXED_COLUMNS, collect_field_columns, write_excel

def test_collect_field_columns_unione_ordinata():
    records = [
        {"campi": {"numero": "1", "data": "x"}},
        {"campi": {"numero": "2", "passeggero": "y"}},
        {"campi": None},
    ]
    assert collect_field_columns(records) == ["data", "numero", "passeggero"]

def test_write_excel_intestazioni_e_valori(tmp_path):
    records = [
        {
            "file": "doc.pdf", "data_email": "2026-05-12", "mittente": "a@b.it",
            "oggetto": "Fattura", "status": "classificato", "fornitore_ok": True,
            "tipo_documento": "fattura", "tipo_label": "Fattura",
            "confidenza_classif": 0.92, "errore": None,
            "campi": {"numero": "ABC123"},
        }
    ]
    out = tmp_path / "ris.xlsx"
    write_excel(records, str(out))

    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == FIXED_COLUMNS + ["numero"]
    row = [c.value for c in ws[2]]
    assert row[headers.index("file")] == "doc.pdf"
    assert row[headers.index("numero")] == "ABC123"
    # None normalizzato a stringa vuota (openpyxl legge le celle vuote come None)
    assert row[headers.index("errore")] in (None, "")
